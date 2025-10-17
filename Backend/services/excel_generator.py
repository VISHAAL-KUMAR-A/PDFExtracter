from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from pathlib import Path
from typing import Dict, Any, List


class ExcelGenerator:
    """Service for generating Excel files from structured data"""

    def generate_excel(self, data: Dict[str, Any], template_id: str, output_path: Path):
        """
        Generate an Excel file from structured data

        Args:
            data: Structured data dictionary
            template_id: Template used for extraction
            output_path: Path where Excel file should be saved
        """
        wb = Workbook()
        # Remove default sheet
        wb.remove(wb.active)

        if template_id == "template1":
            self._generate_template1_excel(wb, data)
        elif template_id == "template2":
            self._generate_template2_excel(wb, data)
        else:
            raise ValueError(f"Unknown template_id: {template_id}")

        wb.save(output_path)

    def _generate_template1_excel(self, wb: Workbook, data: Dict[str, Any]):
        """Generate Excel file for Template 1"""

        # Tab 1: Fund and Investment Vehicle Information
        ws1 = wb.create_sheet("Fund Information")
        self._add_header(ws1, "Fund and Investment Vehicle Information")

        if "fund_information" in data:
            fund_info = data["fund_information"]
            self._add_key_value_section(ws1, fund_info, start_row=3)

        # Tab 2: Fund Manager
        ws2 = wb.create_sheet("Fund Manager")
        self._add_header(ws2, "Fund Manager Details")

        if "fund_manager" in data:
            manager_info = data["fund_manager"]
            self._add_key_value_section(ws2, manager_info, start_row=3)

        # Tab 3: Financial Position
        ws3 = wb.create_sheet("Financial Position")
        self._add_header(ws3, "Fund Investment Vehicle Financial Position")

        if "financial_position" in data:
            financial_info = data["financial_position"]
            self._add_key_value_section(ws3, financial_info, start_row=3)

        # Tab 4: LP Investor Cashflows
        ws4 = wb.create_sheet("LP Investors")
        self._add_header(ws4, "LP Investor Cashflows")

        if "lp_investors" in data and data["lp_investors"]:
            self._add_list_data(ws4, data["lp_investors"], start_row=3)

        # Tab 5: Fund Companies
        ws5 = wb.create_sheet("Fund Companies")
        self._add_header(ws5, "Portfolio Companies")

        if "fund_companies" in data and data["fund_companies"]:
            self._add_list_data(ws5, data["fund_companies"], start_row=3)

        # Tab 6: Company Valuations
        ws6 = wb.create_sheet("Company Valuations")
        self._add_header(ws6, "Company Valuation Details")

        if "company_valuations" in data and data["company_valuations"]:
            self._add_list_data(ws6, data["company_valuations"], start_row=3)

        # Tab 7: Company Financials
        ws7 = wb.create_sheet("Company Financials")
        self._add_header(ws7, "Portfolio Company Financial Information")

        if "company_financials" in data and data["company_financials"]:
            self._add_list_data(ws7, data["company_financials"], start_row=3)

        # Tab 8: Investment History
        ws8 = wb.create_sheet("Investment History")
        self._add_header(ws8, "Historical Transactions")

        if "investment_history" in data and data["investment_history"]:
            self._add_list_data(ws8, data["investment_history"], start_row=3)

    def _generate_template2_excel(self, wb: Workbook, data: Dict[str, Any]):
        """Generate Excel file for Template 2"""

        # Tab 1: Executive Portfolio Summary
        ws1 = wb.create_sheet("Executive Summary")
        self._add_header(ws1, "Executive Portfolio Summary")

        if "executive_summary" in data:
            summary_info = data["executive_summary"]
            self._add_key_value_section(ws1, summary_info, start_row=3)

        # Tab 2: Schedule of Investments
        ws2 = wb.create_sheet("Schedule of Investments")
        self._add_header(ws2, "Schedule of Investments")

        if "schedule_of_investments" in data and data["schedule_of_investments"]:
            self._add_list_data(
                ws2, data["schedule_of_investments"], start_row=3)

        # Tab 3: Statement of Operations
        ws3 = wb.create_sheet("Operations Statement")
        self._add_header(ws3, "Statement of Operations")

        if "operations_statement" in data:
            ops_info = data["operations_statement"]
            self._add_key_value_section(ws3, ops_info, start_row=3)

        # Tab 4: Cashflow Statement
        ws4 = wb.create_sheet("Cashflow Statement")
        self._add_header(ws4, "Statement of Cashflows")

        if "cashflow_statement" in data:
            cashflow_info = data["cashflow_statement"]
            self._add_key_value_section(ws4, cashflow_info, start_row=3)

        # Tab 5: PCAP Statement
        ws5 = wb.create_sheet("PCAP Statement")
        self._add_header(ws5, "Partners' Capital Account (PCAP) Statement")

        if "pcap_statement" in data:
            pcap_info = data["pcap_statement"]
            self._add_key_value_section(ws5, pcap_info, start_row=3)

        # Tab 6: Portfolio Companies Profile
        ws6 = wb.create_sheet("Portfolio Companies")
        self._add_header(ws6, "Portfolio Companies Profile")

        if "portfolio_companies" in data and data["portfolio_companies"]:
            # Convert list fields to strings for display
            companies_data = []
            for company in data["portfolio_companies"]:
                company_copy = company.copy()
                if "key_management" in company_copy and isinstance(company_copy["key_management"], list):
                    company_copy["key_management"] = ", ".join(
                        company_copy["key_management"])
                companies_data.append(company_copy)

            self._add_list_data(ws6, companies_data, start_row=3)

        # Tab 7: Portfolio Companies Financials
        ws7 = wb.create_sheet("Portfolio Financials")
        self._add_header(ws7, "Portfolio Companies Financial Data")

        if "portfolio_financials" in data and data["portfolio_financials"]:
            self._add_list_data(ws7, data["portfolio_financials"], start_row=3)

        # Tab 8: Footnotes
        ws8 = wb.create_sheet("Footnotes")
        self._add_header(ws8, "Footnotes and Disclosures")

        if "footnotes" in data and data["footnotes"]:
            self._add_list_data(ws8, data["footnotes"], start_row=3)

    def _add_header(self, ws, title: str):
        """Add formatted header to worksheet"""
        ws['A1'] = title
        ws['A1'].font = Font(size=14, bold=True, color="FFFFFF")
        ws['A1'].fill = PatternFill(
            start_color="366092", end_color="366092", fill_type="solid")
        ws['A1'].alignment = Alignment(horizontal="center", vertical="center")
        ws.merge_cells('A1:F1')
        ws.row_dimensions[1].height = 25

    def _add_key_value_section(self, ws, data: Dict[str, Any], start_row: int = 1):
        """Add key-value pairs to worksheet"""
        current_row = start_row

        # Header row
        ws.cell(row=current_row, column=1, value="Field")
        ws.cell(row=current_row, column=2, value="Value")

        # Style header
        for col in range(1, 3):
            cell = ws.cell(row=current_row, column=col)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(
                start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
            cell.border = Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin')
            )

        current_row += 1

        # Data rows
        for key, value in data.items():
            # Convert snake_case to Title Case
            field_name = key.replace("_", " ").title()
            ws.cell(row=current_row, column=1, value=field_name)
            ws.cell(row=current_row, column=2,
                    value=value if value is not None else "")

            # Add borders
            for col in range(1, 3):
                cell = ws.cell(row=current_row, column=col)
                cell.border = Border(
                    left=Side(style='thin'),
                    right=Side(style='thin'),
                    top=Side(style='thin'),
                    bottom=Side(style='thin')
                )

            current_row += 1

        # Adjust column widths
        ws.column_dimensions['A'].width = 30
        ws.column_dimensions['B'].width = 40

    def _add_list_data(self, ws, data_list: List[Dict[str, Any]], start_row: int = 1):
        """Add list of dictionaries to worksheet with formatting"""
        if not data_list:
            return

        # Get all unique keys from all dictionaries
        all_keys = []
        seen_keys = set()
        for item in data_list:
            for key in item.keys():
                if key not in seen_keys:
                    all_keys.append(key)
                    seen_keys.add(key)

        # Add headers
        for col_idx, column_name in enumerate(all_keys, 1):
            cell = ws.cell(row=start_row, column=col_idx)
            # Convert snake_case to Title Case
            cell.value = column_name.replace("_", " ").title()
            cell.font = Font(bold=True)
            cell.fill = PatternFill(
                start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
            cell.border = Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin')
            )

        # Add data rows
        for row_idx, item in enumerate(data_list, start_row + 1):
            for col_idx, key in enumerate(all_keys, 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                value = item.get(key)
                cell.value = value if value is not None else ""
                cell.border = Border(
                    left=Side(style='thin'),
                    right=Side(style='thin'),
                    top=Side(style='thin'),
                    bottom=Side(style='thin')
                )

        # Adjust column widths
        for col_idx in range(1, len(all_keys) + 1):
            # Use appropriate letter for column (A, B, C, ... Z, AA, AB, etc.)
            if col_idx <= 26:
                col_letter = chr(64 + col_idx)
            else:
                col_letter = chr(64 + (col_idx - 1) // 26) + \
                    chr(65 + (col_idx - 1) % 26)
            ws.column_dimensions[col_letter].width = 20
